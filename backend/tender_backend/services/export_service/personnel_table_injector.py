from __future__ import annotations

import re
from uuid import UUID

from docx.document import Document
from psycopg import Connection

from tender_backend.services.export_service.personnel_table_renderer import (
    _COLUMNS,
    PersonnelTableRenderer,
)


_ANCHOR_RE = re.compile(r"\{\{personnel_table\}\}")


class PersonnelTableInjector:
    def __init__(self, document: Document, conn: Connection, *, project_id: UUID):
        self._document = document
        self._conn = conn
        self._project_id = project_id
        self._renderer = PersonnelTableRenderer()

    def inject_all(self) -> int:
        rows = self._renderer.render_personnel_preview(self._conn, project_id=self._project_id)
        count = 0
        for paragraph in list(self._document.paragraphs):
            if not _ANCHOR_RE.search(paragraph.text or ""):
                continue
            table = self._build_table(rows=rows)
            paragraph._element.addnext(table._element)  # noqa: SLF001
            paragraph._element.getparent().remove(paragraph._element)  # noqa: SLF001
            count += 1
        return count

    def _build_table(self, *, rows: list[dict[str, str]]):
        columns = [label for label, _source in _COLUMNS]
        data_rows = rows if rows else [{columns[0]: "无", **{label: "" for label in columns[1:]}}]
        table = self._document.add_table(rows=len(data_rows) + 1, cols=len(columns))
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        for index, label in enumerate(columns):
            table.cell(0, index).text = label
        for row_index, row in enumerate(data_rows, start=1):
            for column_index, label in enumerate(columns):
                table.cell(row_index, column_index).text = str(row.get(label, ""))
        return table
