from __future__ import annotations

import re
from uuid import UUID

from docx.document import Document
from psycopg import Connection

from tender_backend.services.export_service.equipment_table_renderer import (
    _ASSET_LABELS,
    _COLUMNS,
    EquipmentTableRenderer,
)


_ANCHOR_RE = re.compile(r"\{\{equipment_table:(vehicle|machine|tool|safety)\}\}")


class EquipmentTableInjector:
    def __init__(self, document: Document, conn: Connection, *, project_id: UUID):
        self._document = document
        self._conn = conn
        self._project_id = project_id
        self._renderer = EquipmentTableRenderer()

    def inject_all(self) -> int:
        preview = self._renderer.render_equipment_preview(self._conn, project_id=self._project_id)
        count = 0
        for paragraph in list(self._document.paragraphs):
            match = _ANCHOR_RE.search(paragraph.text or "")
            if not match:
                continue
            asset_type = match.group(1)
            table = self._build_table(asset_type=asset_type, rows=preview.get(asset_type, []))
            paragraph._element.addnext(table._element)  # noqa: SLF001
            paragraph._element.getparent().remove(paragraph._element)  # noqa: SLF001
            count += 1
        return count

    def _build_table(self, *, asset_type: str, rows: list[dict[str, str]]):
        columns = [label for label, _source in _COLUMNS[asset_type]]
        data_rows = rows if rows else [{columns[0]: "无", **{label: "" for label in columns[1:]}}]
        table = self._document.add_table(rows=len(data_rows) + 1, cols=len(columns))
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        for index, label in enumerate(columns):
            table.cell(0, index).text = label
        for row_index, row in enumerate(data_rows, start=1):
            for column_index, label in enumerate(columns):
                table.cell(row_index, column_index).text = str(row.get(label, ""))
        return table
