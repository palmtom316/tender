from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document

from scripts.build_sgcc_distribution_business_full_template import BUSINESS_PACKAGE_KEY, ALL, build


def _text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def test_business_full_template_builder_uses_single_docx_placeholders_without_internal_notes(tmp_path: Path) -> None:
    output = tmp_path / "商务标完整模板.docx"

    build(output)

    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert [name for name in archive.namelist() if name.startswith("word/media/")] == []
    text = _text(output)
    assert BUSINESS_PACKAGE_KEY == "sgcc_distribution_business_full_v1"
    assert len(ALL) == 60
    assert "1. 商务偏差表" in text
    assert "（本页不编辑正文）" in text
    assert "{{ business_deviation_rows }}" in text
    assert "{{ asset:business_license:1 }}" in text
    assert "{{ asset:financial_statement:n }}" in text
    assert "{{ asset:bid_bond_assets:n }}" in text
    assert "【模板用途】" not in text
    assert "生成正式投标文件时" not in text
